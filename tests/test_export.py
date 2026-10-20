"""T-22：三格式导出共用同一份 payload，内容必须一致。"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from cause_trace.export import render, render_excel, render_markdown, render_pdf
from cause_trace.report import SECTION_DEFINITIONS, empty_section_payload

META = {
    "title": "8 月 GMV 下滑归因报告",
    "domain_id": 1,
    "metric_id": 1,
    "window_start": "2026-08-12",
    "window_end": "2026-08-18",
    "baseline": "dod",
}

CONTENTS = {
    "background": {"summary": "GMV 日环比 -12.3%，差异额 -1,432 万"},
    "factor_contribution": {
        "contributions": [
            {"factor": "访客数", "contribution": -320.0, "contribution_rate": -0.45, "direction": "down"},
            {"factor": "转化率", "contribution": -228.0, "contribution_rate": -0.32, "direction": "down"},
        ]
    },
    "action_plan": {
        "actions": [
            {"factor": "渠道预算重配", "role": "渠道运营负责人", "period_days": 7, "expected_recovery": 286.0}
        ]
    },
}


def _payload() -> dict:
    return empty_section_payload(META, CONTENTS)


def test_markdown_contains_all_sections():
    artifact = render_markdown(_payload())
    text = artifact.content.decode("utf-8")

    assert artifact.filename.endswith(".md")
    for index, (_, title) in enumerate(SECTION_DEFINITIONS, start=1):
        assert f"## {index}. {title}" in text
    assert "-1,432 万" in text


def test_markdown_renders_contribution_table():
    text = render_markdown(_payload()).content.decode("utf-8")
    assert "| 因子 | 贡献量 | 贡献率 | 方向 |" in text
    assert "| 访客数 | -320.0 | -45.0% | down |" in text


def test_markdown_renders_action_table():
    text = render_markdown(_payload()).content.decode("utf-8")
    assert "| 渠道预算重配 | 渠道运营负责人 | 7 | 286.0 |" in text


def test_excel_has_one_sheet_per_section_plus_overview():
    artifact = render_excel(_payload())
    workbook = load_workbook(io.BytesIO(artifact.content))

    assert workbook.sheetnames[0] == "报告概览"
    assert len(workbook.sheetnames) == 8
    assert workbook["报告概览"]["A1"].value == "报告标题"
    assert workbook["2 因子拆解"]["A1"].value == "因子贡献率拆解"


def test_excel_writes_contribution_rows():
    workbook = load_workbook(io.BytesIO(render_excel(_payload()).content))
    sheet = workbook["2 因子拆解"]
    values = [cell.value for cell in sheet[3]]
    assert values == ["factor", "contribution", "contribution_rate", "direction"]


def test_invalid_payload_is_rejected_before_export():
    payload = _payload()
    payload["sections"].pop()
    with pytest.raises(Exception) as exc_info:
        render_markdown(payload)
    assert "payload" in str(exc_info.value) or "章节" in str(exc_info.value)


def _weasyprint_available() -> bool:
    try:
        import weasyprint  # noqa: F401
    except Exception:  # 缺 Pango/Cairo 时 import 会抛 OSError
        return False
    return True


requires_weasyprint = pytest.mark.skipif(
    not _weasyprint_available(),
    reason="本地无 Pango/Cairo，PDF 渲染在容器内验证（scripts/e2e_pdf_container.sh）",
)


@requires_weasyprint
def test_pdf_has_valid_header_and_non_trivial_size():
    artifact = render_pdf(_payload())
    assert artifact.content.startswith(b"%PDF-")
    assert artifact.content.rstrip().endswith(b"%%EOF")
    # WeasyPrint 带 CSS 与字体的输出远大于纯文本 PDF（后者约 1.5KB）
    assert len(artifact.content) > 10_000


@requires_weasyprint
def test_pdf_renders_chinese_and_pagination():
    """WeasyPrint 会内嵌字体子集，中文与页脚分页标记应出现在 PDF 里。"""
    artifact = render_pdf(_payload())
    assert b"/FontFile2" in artifact.content or b"/FontFile3" in artifact.content
    # 至少一页
    assert b"/Type /Page" in artifact.content


def test_pdf_template_contains_css_and_svg():
    """模板本身必须带 @page 分页规则与内联 SVG 图表，不依赖容器即可断言。"""
    from cause_trace.export.pdf_template import build_report_html
    from cause_trace.report import empty_section_payload

    payload = empty_section_payload(
        {
            "title": "模板测试",
            "domain_id": 1,
            "metric_id": 1,
            "window_start": "2026-08-12",
            "window_end": "2026-08-18",
            "baseline": "dod",
        },
        {
            "factor_contribution": {
                "contributions": [
                    {"factor": "抖音", "contribution": -100.0, "contribution_rate": -0.5, "direction": "down"}
                ]
            }
        },
    )
    html = build_report_html(payload)
    assert "@page" in html
    assert "counter(page)" in html
    assert "<svg" in html


def test_render_dispatches_all_three_formats():
    expected_extension = {"markdown": ".md", "excel": ".xlsx", "pdf": ".pdf"}
    for name, extension in expected_extension.items():
        if name == "pdf" and not _weasyprint_available():
            continue
        artifact = render(name, _payload())
        assert artifact.content
        assert artifact.filename.endswith(extension)


def test_unknown_format_is_rejected():
    with pytest.raises(ValueError):
        render("docx", _payload())


def test_html_template_escapes_user_input():
    """用户输入进 HTML 前必须转义，防止模板注入（Spec 16.4）。"""
    from cause_trace.export.pdf_template import build_report_html
    from cause_trace.report import empty_section_payload

    payload = empty_section_payload(
        {
            "title": "<script>alert(1)</script>",
            "domain_id": 1,
            "metric_id": 1,
            "window_start": "2026-08-12",
            "window_end": "2026-08-18",
            "baseline": "dod",
        },
        {"background": {"summary": "<img src=x onerror=alert(1)>"}},
    )
    html = build_report_html(payload)
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html
    assert "<img src=x" not in html
